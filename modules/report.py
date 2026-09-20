"""Word report generation with python-docx."""

from __future__ import annotations

from datetime import datetime
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path
from typing import Mapping

import pandas as pd
from docx import Document
from docx.shared import Inches

from modules.utils import ensure_parent


VERSION_PACKAGES = (
    "rasterio",
    "numpy",
    "pandas",
    "matplotlib",
    "geopandas",
    "shapely",
    "pyproj",
    "openpyxl",
    "python-docx",
)


def generate_word_report(
    output_path: Path,
    title: str,
    summary: Mapping[str, object],
    tables: Mapping[str, pd.DataFrame],
    generated_files: Mapping[str, Path],
) -> Path:
    """Create a DOCX report with parameters, tables, figures and file list."""

    ensure_parent(output_path)
    document = Document()
    document.add_heading(title, level=0)
    document.add_paragraph(
        "Автоматически сформированный отчет программного комплекса NDVI Monitor."
    )
    document.add_paragraph(
        f"Дата и время обработки: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )

    document.add_heading("Параметры исследования", level=1)
    _add_mapping_table(document, summary)

    document.add_heading("Версии библиотек", level=1)
    _add_mapping_table(document, _package_versions())

    for name, table in tables.items():
        document.add_heading(name, level=1)
        _add_dataframe_table(document, table)

    figure_paths = [
        path
        for path in generated_files.values()
        if Path(path).suffix.lower() == ".png" and Path(path).exists()
    ]
    if figure_paths:
        document.add_heading("Рисунки", level=1)
        for figure_path in figure_paths:
            document.add_paragraph(Path(figure_path).name)
            document.add_picture(str(figure_path), width=Inches(6.2))

    document.add_heading("Созданные файлы", level=1)
    _add_mapping_table(document, {key: str(path) for key, path in generated_files.items()})

    document.save(output_path)
    return output_path


def _add_mapping_table(document: Document, values: Mapping[str, object]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Параметр"
    header[1].text = "Значение"
    for key, value in values.items():
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = _format_cell(value)


def _add_dataframe_table(document: Document, dataframe: pd.DataFrame) -> None:
    if dataframe.empty:
        document.add_paragraph("Нет данных для отображения.")
        return

    display_frame = dataframe.drop(columns=["geometry"], errors="ignore")
    table = document.add_table(rows=1, cols=len(display_frame.columns))
    table.style = "Table Grid"
    for index, column in enumerate(display_frame.columns):
        table.rows[0].cells[index].text = str(column)

    for _, row in display_frame.iterrows():
        cells = table.add_row().cells
        for index, column in enumerate(display_frame.columns):
            cells[index].text = _format_cell(row[column])


def _format_cell(value: object) -> str:
    if isinstance(value, float):
        return f"{value:.6g}"
    return str(value)


def _package_versions() -> dict[str, str]:
    versions = {}
    for package_name in VERSION_PACKAGES:
        try:
            versions[package_name] = version(package_name)
        except PackageNotFoundError:
            versions[package_name] = "not installed"
    return versions
