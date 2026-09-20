"""End-to-end verification utility for NDVI Monitor."""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import numpy as np
import pandas as pd
import rasterio
from docx import Document

from config import (
    BASE_YEAR,
    CLASS_NODATA,
    FORECAST_YEAR,
    NDVI_MAX,
    NDVI_MIN,
    ProjectPaths,
    TARGET_YEAR,
)
from modules.loader import validate_inputs
from modules.pipeline import MonitorPipeline
from modules.utils import assert_same_grid, read_masked_raster


@dataclass(frozen=True)
class CheckResult:
    """Result of one verification check."""

    name: str
    passed: bool
    message: str


def build_parser() -> argparse.ArgumentParser:
    """Create command line parser for verification."""

    parser = argparse.ArgumentParser(description="Verify NDVI Monitor pipeline outputs.")
    parser.add_argument("--input", type=Path, default=None, help="Input data directory.")
    parser.add_argument(
        "--skip-run",
        action="store_true",
        help="Do not run the pipeline before checking outputs.",
    )
    return parser


def main() -> int:
    """Run verification and print a PASS/FAIL report."""

    _configure_console_encoding()
    args = build_parser().parse_args()
    paths = ProjectPaths.from_base()
    input_dir = args.input or paths.input_dir
    context: dict[str, object] = {"paths": paths, "input_dir": input_dir}

    results: list[CheckResult] = []
    results.append(_run_check("input_validation", lambda: _check_inputs(input_dir, context)))

    if not args.skip_run:
        results.append(_run_check("pipeline_run", lambda: _run_pipeline(paths, input_dir, context)))

    checks: list[tuple[str, Callable[[], str]]] = [
        ("export_files", lambda: _check_export_files(paths)),
        ("spatial_grid", lambda: _check_spatial_grid(paths)),
        ("ndvi_range", lambda: _check_raster_range(paths.rasters_dir / f"ndvi_{BASE_YEAR}_valid.tif", NDVI_MIN, NDVI_MAX)),
        ("target_ndvi_range", lambda: _check_raster_range(paths.rasters_dir / f"ndvi_{TARGET_YEAR}_valid.tif", NDVI_MIN, NDVI_MAX)),
        ("delta_range", lambda: _check_raster_range(paths.rasters_dir / "delta_ndvi.tif", -2.0, 2.0)),
        ("forecast_range", lambda: _check_raster_range(paths.rasters_dir / f"ndvi_forecast_{FORECAST_YEAR}.tif", NDVI_MIN, NDVI_MAX)),
        ("nodata_masks", lambda: _check_nodata_masks(paths)),
        ("statistics", lambda: _check_statistics(paths)),
        ("transition_matrix", lambda: _check_transition_matrix(paths)),
        ("change_polygons", lambda: _check_change_polygons(paths)),
        ("word_report", lambda: _check_word_report(paths)),
        ("png_outputs", lambda: _check_png_outputs(paths)),
    ]
    for name, check in checks:
        results.append(_run_check(name, check))

    for result in results:
        status = "PASS" if result.passed else "FAIL"
        print(f"{status:4} | {result.name} | {result.message}")

    failed = [result for result in results if not result.passed]
    if failed:
        print(f"\nИтог: FAIL ({len(failed)} проверок не пройдено)")
        return 1
    print("\nИтог: PASS (все проверки пройдены)")
    return 0


def _check_inputs(input_dir: Path, context: dict[str, object]) -> str:
    bundle = validate_inputs(input_dir)
    context["input_bundle"] = bundle
    return (
        f"variant={bundle.variant.value}, rasters={len(bundle.rasters)}, "
        f"boundary={bundle.boundary.name}"
    )


def _run_pipeline(
    paths: ProjectPaths,
    input_dir: Path,
    context: dict[str, object],
) -> str:
    result = MonitorPipeline(paths).run(input_dir=input_dir)
    context["pipeline_result"] = result
    return f"generated_files={len(result.generated_files)}, tables={len(result.tables)}"


def _check_export_files(paths: ProjectPaths) -> str:
    required = [
        paths.rasters_dir / f"ndvi_{BASE_YEAR}_valid.tif",
        paths.rasters_dir / f"ndvi_{TARGET_YEAR}_valid.tif",
        paths.rasters_dir / "common_valid_mask.tif",
        paths.rasters_dir / "delta_ndvi.tif",
        paths.rasters_dir / "delta_class.tif",
        paths.rasters_dir / f"ndvi_forecast_{FORECAST_YEAR}.tif",
        paths.tables_dir / "ndvi_monitor_results.xlsx",
        paths.tables_dir / "transition_matrix.xlsx",
        paths.tables_dir / "largest_changes.xlsx",
        paths.rasters_dir / "change_polygons.gpkg",
        paths.reports_dir / "ndvi_monitor_report.docx",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise AssertionError("Missing outputs: " + ", ".join(missing))
    return f"checked={len(required)}"


def _check_spatial_grid(paths: ProjectPaths) -> str:
    rasters = [
        paths.rasters_dir / f"ndvi_{BASE_YEAR}_valid.tif",
        paths.rasters_dir / f"ndvi_{TARGET_YEAR}_valid.tif",
        paths.rasters_dir / "delta_ndvi.tif",
        paths.rasters_dir / "delta_class.tif",
        paths.rasters_dir / f"ndvi_forecast_{FORECAST_YEAR}.tif",
    ]
    reference = rasters[0]
    for raster_path in rasters[1:]:
        assert_same_grid(reference, raster_path)
    with rasterio.open(reference) as dataset:
        return (
            f"crs={dataset.crs}, size={dataset.width}x{dataset.height}, "
            f"transform={dataset.transform}"
        )


def _check_raster_range(path: Path, minimum: float, maximum: float) -> str:
    values, _ = read_masked_raster(path)
    valid = ~np.ma.getmaskarray(values) & np.isfinite(values.data)
    if not valid.any():
        raise AssertionError(f"No valid pixels in {path}")
    actual_min = float(np.min(values.data[valid]))
    actual_max = float(np.max(values.data[valid]))
    if actual_min < minimum or actual_max > maximum:
        raise AssertionError(
            f"{path.name} range {actual_min}..{actual_max} is outside "
            f"{minimum}..{maximum}"
        )
    return f"{path.name}: {actual_min:.6g}..{actual_max:.6g}"


def _check_nodata_masks(paths: ProjectPaths) -> str:
    checked = 0
    for raster_path in paths.rasters_dir.glob("*.tif"):
        with rasterio.open(raster_path) as dataset:
            if dataset.nodata is None:
                raise AssertionError(f"No NoData metadata in {raster_path}")
            data = dataset.read(1)
            mask = dataset.dataset_mask()
            if np.issubdtype(data.dtype, np.floating):
                nodata_pixels = np.isclose(data, float(dataset.nodata))
            else:
                nodata_pixels = data == dataset.nodata
            if nodata_pixels.any() and np.any(mask[nodata_pixels] != 0):
                raise AssertionError(f"NoData mask is invalid in {raster_path}")
            checked += 1
    if checked == 0:
        raise AssertionError("No GeoTIFF outputs found for NoData verification.")
    return f"checked_rasters={checked}"


def _check_statistics(paths: ProjectPaths) -> str:
    workbook = paths.tables_dir / "ndvi_monitor_results.xlsx"
    sheets = pd.read_excel(workbook, sheet_name=None)
    statistic_sheets = [
        name for name in sheets if name.startswith("ndvi_statistics") or name == "delta_statistics"
    ]
    if not statistic_sheets:
        raise AssertionError("No statistics sheets found.")
    for sheet_name in statistic_sheets:
        columns = set(sheets[sheet_name].columns)
        required = {"mean", "minimum", "maximum", "median", "std", "pixel_count"}
        if not required.issubset(columns):
            raise AssertionError(f"Statistics sheet lacks columns: {sheet_name}")
    return f"statistic_sheets={len(statistic_sheets)}"


def _check_transition_matrix(paths: ProjectPaths) -> str:
    table = pd.read_excel(paths.tables_dir / "transition_matrix.xlsx")
    required = {"from_class", "to_class", "pixel_count", "area_m2", "percent"}
    if not required.issubset(set(table.columns)):
        raise AssertionError("Transition matrix lacks required columns.")
    class_2015 = paths.rasters_dir / f"ndvi_classes_{BASE_YEAR}.tif"
    class_2025 = paths.rasters_dir / f"ndvi_classes_{TARGET_YEAR}.tif"
    with rasterio.open(class_2015) as first, rasterio.open(class_2025) as second:
        first_classes = first.read(1)
        second_classes = second.read(1)
    expected_pixels = int(
        np.count_nonzero((first_classes != CLASS_NODATA) & (second_classes != CLASS_NODATA))
    )
    actual_pixels = int(table["pixel_count"].sum())
    if expected_pixels != actual_pixels:
        raise AssertionError(
            f"Transition pixel mismatch: expected={expected_pixels}, actual={actual_pixels}"
        )
    percent_sum = float(table["percent"].sum())
    if expected_pixels and not np.isclose(percent_sum, 100.0, atol=0.01):
        raise AssertionError(f"Transition percent sum is {percent_sum}")
    return f"pixels={actual_pixels}, percent_sum={percent_sum:.4f}"


def _check_change_polygons(paths: ProjectPaths) -> str:
    vector_path = paths.rasters_dir / "change_polygons.gpkg"
    excel_path = paths.tables_dir / "largest_changes.xlsx"
    import geopandas as gpd

    vector = gpd.read_file(vector_path)
    table = pd.read_excel(excel_path)
    required = {
        "change_type",
        "area_m2",
        "area_ha",
        "bbox_min_x",
        "bbox_min_y",
        "bbox_max_x",
        "bbox_max_y",
        "centroid_lon",
        "centroid_lat",
    }
    if not required.issubset(set(table.columns)):
        raise AssertionError("Largest changes table lacks required columns.")
    return f"vector_features={len(vector)}, excel_rows={len(table)}"


def _check_word_report(paths: ProjectPaths) -> str:
    document = Document(paths.reports_dir / "ndvi_monitor_report.docx")
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    if not any("NDVI Monitor" in text for text in paragraphs):
        raise AssertionError("Word report title was not found.")
    if len(document.tables) == 0:
        raise AssertionError("Word report has no tables.")
    if len(document.inline_shapes) == 0:
        raise AssertionError("Word report has no embedded figures.")
    return (
        f"paragraphs={len(paragraphs)}, tables={len(document.tables)}, "
        f"figures={len(document.inline_shapes)}"
    )


def _check_png_outputs(paths: ProjectPaths) -> str:
    png_files = list(paths.figures_dir.glob("*.png"))
    if not png_files:
        raise AssertionError("No PNG outputs found.")
    empty = [str(path) for path in png_files if path.stat().st_size == 0]
    if empty:
        raise AssertionError("Empty PNG files: " + ", ".join(empty))
    return f"png_files={len(png_files)}"


def _run_check(name: str, check: Callable[[], str]) -> CheckResult:
    try:
        message = check()
        return CheckResult(name, True, message)
    except Exception as exc:
        return CheckResult(name, False, str(exc))


def _configure_console_encoding() -> None:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")


if __name__ == "__main__":
    raise SystemExit(main())
